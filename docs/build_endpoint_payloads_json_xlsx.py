#!/usr/bin/env python3
"""Build one Excel workbook: eight tabs, each the full JSON payload line-by-line."""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "SAT-Datashare-Kenya-Endpoint-Payloads-JSON.xlsx"

PAYLOAD_JSON = {
    "Invoice": ROOT / "payloads" / "invoice.json",
    "Order": ROOT / "payloads" / "order.json",
    "Product": ROOT / "payloads" / "product.json",
    "Route": ROOT / "payloads" / "route.json",
    "SalesRep": ROOT / "payloads" / "salesrep.json",
    "Customer": ROOT / "payloads" / "customer.json",
    "Stock": ROOT / "payloads" / "stock.json",
    "Warehouse": ROOT / "payloads" / "warehouse.json",
}

SHEET_ORDER = [
    "Invoice",
    "Order",
    "Product",
    "Route",
    "SalesRep",
    "Customer",
    "Stock",
    "Warehouse",
]

# Blue sheet tabs, matching the tabular workbook.
TAB_COLORS = {
    "Invoice": "1B365D",
    "Order": "1E3D63",
    "Product": "214269",
    "Route": "24476F",
    "SalesRep": "274C75",
    "Customer": "2A517B",
    "Stock": "2D5681",
    "Warehouse": "305B87",
}

FONT_NAME = "Consolas"
FONT_SIZE = 11

KEY_FONT = InlineFont(rFont=FONT_NAME, color="A31515", sz=FONT_SIZE)
STR_FONT = InlineFont(rFont=FONT_NAME, color="0000FF", sz=FONT_SIZE)
NUM_FONT = InlineFont(rFont=FONT_NAME, color="098658", sz=FONT_SIZE)
LIT_FONT = InlineFont(rFont=FONT_NAME, color="0000FF", sz=FONT_SIZE)
PUNCT_FONT = InlineFont(rFont=FONT_NAME, color="000000", sz=FONT_SIZE)
WS_FONT = InlineFont(rFont=FONT_NAME, color="000000", sz=FONT_SIZE)

PLAIN = Font(name=FONT_NAME, size=FONT_SIZE, color="000000")
LEFT = Alignment(horizontal="left", vertical="center")
WHITE = PatternFill("solid", fgColor="FFFFFF")

TOKEN = re.compile(
    r"""
    (?P<ws>\s+)
    | (?P<lit>true|false|null)
    | (?P<num>-?\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)
    | (?P<str>"(?:\\.|[^"\\])*")(?P<colon>\s*:)?
    | (?P<punct>[{\[\]},:])
    """,
    re.VERBOSE,
)


def colorize_line(line: str) -> CellRichText:
    blocks: list[TextBlock] = []
    pos = 0
    for match in TOKEN.finditer(line):
        if match.start() > pos:
            blocks.append(TextBlock(PUNCT_FONT, line[pos : match.start()]))
        if match.group("ws") is not None:
            blocks.append(TextBlock(WS_FONT, match.group("ws")))
        elif match.group("lit") is not None:
            blocks.append(TextBlock(LIT_FONT, match.group("lit")))
        elif match.group("num") is not None:
            blocks.append(TextBlock(NUM_FONT, match.group("num")))
        elif match.group("str") is not None:
            blocks.append(TextBlock(KEY_FONT if match.group("colon") else STR_FONT, match.group("str")))
            if match.group("colon"):
                blocks.append(TextBlock(PUNCT_FONT, match.group("colon")))
        elif match.group("punct") is not None:
            blocks.append(TextBlock(PUNCT_FONT, match.group("punct")))
        pos = match.end()
    if pos < len(line):
        blocks.append(TextBlock(PUNCT_FONT, line[pos:]))
    if not blocks:
        return CellRichText(TextBlock(PUNCT_FONT, line or " "))
    return CellRichText(*blocks)


def write_json_sheet(wb: Workbook, name: str, payload: dict) -> int:
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = TAB_COLORS[name]
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.6)
    ws.oddHeader.left.text = f"SAT Datashare · {name} JSON"
    ws.oddFooter.left.text = "Confidential · v1.0"
    ws.oddFooter.right.text = "Page &P"
    ws.sheet_view.showGridLines = True

    text = json.dumps(payload, indent=4, ensure_ascii=False)
    lines = text.splitlines()
    for i, line in enumerate(lines, 1):
        cell = ws.cell(i, 1)
        cell.value = colorize_line(line)
        cell.alignment = LEFT
        cell.fill = WHITE
        ws.row_dimensions[i].height = 15
    ws.column_dimensions["A"].width = 110
    return len(lines)


def main() -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    wb.properties.creator = "SAT Datashare"
    wb.properties.lastModifiedBy = "SAT Datashare"
    wb.properties.title = "SAT Datashare Kenya endpoint JSON payloads"
    wb.properties.subject = "Full JSON payload per endpoint, one tab each"
    wb.properties.keywords = "SAT Datashare, Kenya, JSON"

    for name in SHEET_ORDER:
        path = PAYLOAD_JSON[name]
        payload = json.loads(path.read_text(encoding="utf-8"))
        rows = payload.get("data") or []
        if not rows:
            raise SystemExit(f"{path} has no data array")
        n_lines = write_json_sheet(wb, name, payload)
        print(f"  {name}: {len(rows)} records, {n_lines} JSON lines")

    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
