#!/usr/bin/env python3
"""Client-facing SAT Datashare Experience API test guide. No PRC / SYS detail."""
from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "SAT-Datashare-Client-Test-Guide-Kenya.docx"

NAVY = RGBColor(0x1F, 0x4E, 0x79)
TEAL = RGBColor(0x2E, 0x75, 0xB6)
TEXT = RGBColor(0x2D, 0x2D, 0x2D)
MUTED = RGBColor(0x5B, 0x65, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY_HEX = "1F4E79"
ROW_HEX = "E9F0F7"
CODE_HEX = "F3F4F6"
AMBER_HEX = "FFF4D6"
AMBER_BORDER = "E0C36A"

VERSION = "v1.1"
VERSION_DATE = "25 September 2026"
AUTHOR = "SAT Datashare"
PAGE_SIZE = 10000

BASE = "https://exp-sat-datashare-prod-api-vx7q2k.2ky31l-1.deu-c1.eu1.cloudhub.io"
TENANT_ID = "f4faf003-d90b-4832-9df8-c8a22d29bdd4"
CLIENT_ID = "de772600-0d1f-4492-90b0-a57fcc284cf9"
TOKEN_URL = f"https://login.microsoftonline.com/{TENANT_ID}/oauth2/v2.0/token"
SCOPE = "https://graph.microsoft.com/.default"
CONTENT_WIDTH_CM = 17.4
TWIPS_PER_CM = 567

SAMPLES = {
    "Invoice": {
        "STOREID": 2655,
        "Discount": 0,
        "SalesRepID": 111,
        "VATAmount": "80.46",
        "RouteID": "",
        "DistributorID": "1000097205",
        "ProductID": "3286535",
        "WarehouseID": 1925,
        "OrderNumber": 265,
        "InvoiceNumber": 2857,
        "InvoiceDate": "24/09/2026",
        "InvoiceTime": "15:46:26",
        "InvoiceStatus": "Delivered",
        "UnitOfMeasure": "PIECE",
        "Currency": "KES",
        "QuantityInvoiced": 10,
        "AmountInvoiced": "583.30",
        "TransType": "Sale",
    },
    "Order": {
        "StoreID": 2348,
        "SalesRepID": 108,
        "RouteID": "",
        "ProductID": "3049801",
        "DistributorID": "1000097205",
        "WarehouseID": 2503,
        "OrderNumber": 320,
        "OrderDate": "24/09/2026",
        "OrderStartTime": "",
        "OrderEndTime": "",
        "OrderStatus": "Open",
        "UnitOfMeasure": "PIECE",
        "Currency": "KES",
        "QuantityOrdered": 5,
        "AmountOrdered": "590.90",
        "Discount": 0,
        "VATAmount": "81.50",
    },
    "Product": {
        "ProductID": "3268896",
        "ProductSKU": "AWICK FRESHMATIC ROSE  + GADGET SEEDING PRICE 250ML (4)",
        "UnitofMeasure": "PIECE",
    },
    "Route": {
        "RouteID": 110,
        "RouteName": "MOSES RB TUESDAY",
        "SalesRepID": 110,
        "RouteType": "Weekly",
        "SalesRepType": "Van Sales",
        "Date": "07/09/2026",
    },
    "SalesRep": {
        "SalesRepID": 100,
        "SalesRepName": "ISAAC MUGE",
        "DistributorID": "1000097205",
        "WarehouseID": 183,
        "SalesRepType": "Admin",
    },
    "Customer": {
        "StoreID": 1980,
        "StoreName": "ONE ZERO ONE SELFRIGES",
        "StoreCategory": "SHOP AND BROWSER",
        "Channel": "GENERAL TRADE",
        "StoreClassification": "GENERAL TRADE",
        "WarehouseID": None,
        "Longitude": 37.5852265,
        "Latitude": 0.3526255,
        "City": "MERU",
        "Region": "MOUNTAIN",
        "County": "",
        "StoreStatus": 1,
        "StoreCreationDate": "10/09/2026",
        "StoreSize": "",
    },
    "Stock": {
        "ProductSKU": "3300090",
        "UnitofMeasure": "PIECE",
        "DistributorWarehouseCode": "EMBU STOCKPOINT",
        "Date": "22/09/2026",
        "InventoryQuantity": 21,
        "InventoryPrice": None,
    },
    "Warehouse": {
        "WarehouseID": 183,
        "WarehouseName": "GIKAMBURA STOCKPOINT",
        "DistributorID": "1000097205",
    },
}

DATE_PARAMS = [
    ("fromDate", "Yes", "YYYY-MM-DD", "2026-09-24", "Inclusive start date. Date only. Do not send a time."),
    ("toDate", "Yes", "YYYY-MM-DD", "2026-09-24", "Inclusive end date. Date only. Do not send a time."),
    ("pageNumber", "No", "Integer, 0-based", "0", "Omit on the first call. Use 1, 2, … if a next page is returned."),
]
PAGE_PARAMS = [
    ("pageNumber", "No", "Integer, 0-based", "0", "Omit on the first call. Use 1, 2, … if a next page is returned."),
]

ENDPOINTS = [
    {
        "name": "Invoice",
        "path": "/api/v1/Invoice",
        "params": DATE_PARAMS,
        "query": "fromDate=2026-09-24&toDate=2026-09-24&pageNumber=0",
        "notes": "Returns invoiced sale lines for the date range.",
    },
    {
        "name": "Order",
        "path": "/api/v1/Order",
        "params": DATE_PARAMS,
        "query": "fromDate=2026-09-24&toDate=2026-09-24&pageNumber=0",
        "notes": "Returns order lines for the date range.",
    },
    {
        "name": "Product",
        "path": "/api/v1/Product",
        "params": [],
        "query": "",
        "notes": "Returns the product catalogue.",
    },
    {
        "name": "Route",
        "path": "/api/v1/Route",
        "params": [],
        "query": "",
        "notes": "Returns sales routes.",
    },
    {
        "name": "SalesRep",
        "path": "/api/v1/SalesRep",
        "params": [],
        "query": "",
        "notes": "Returns sales representatives.",
    },
    {
        "name": "Customer",
        "path": "/api/v1/Customer",
        "params": DATE_PARAMS,
        "query": "fromDate=2026-09-24&toDate=2026-09-24&pageNumber=0",
        "notes": "Returns stores created in the date range.",
    },
    {
        "name": "Stock",
        "path": "/api/v1/Stock",
        "params": PAGE_PARAMS,
        "query": "pageNumber=0",
        "notes": "Returns inventory by warehouse.",
    },
    {
        "name": "Warehouse",
        "path": "/api/v1/Warehouse",
        "params": [],
        "query": "",
        "notes": "Returns warehouses.",
    },
]


def shade(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    existing = tcPr.find(qn("w:shd"))
    if existing is not None:
        tcPr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, color="BFBFBF") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    existing = tcPr.find(qn("w:tcBorders"))
    if existing is not None:
        tcPr.remove(existing)
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_run(run, *, name="Calibri", size=11, bold=False, color=TEXT, italic=False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = name
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:eastAsia"), name)


def write_cell(cell, text, *, bold=False, size=10, color=TEXT, font="Calibri", fill=None, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if align:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, name=font, size=size, bold=bold, color=color)
    if fill:
        shade(cell, fill)
    set_cell_border(cell)


def add_para(
    doc,
    text,
    *,
    size=11,
    bold=False,
    color=TEXT,
    space_after=8,
    space_before=0,
    align=None,
    keep_with_next=False,
):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = keep_with_next
    if align:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    return p


def configure_heading_styles(doc) -> None:
    for style_name, size, before, after in (
        ("Heading 1", 16, 16, 8),
        ("Heading 2", 13, 14, 6),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY
        style.font.italic = False
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:ascii"), "Calibri")
        rFonts.set(qn("w:hAnsi"), "Calibri")
        rFonts.set(qn("w:cs"), "Calibri")
        rFonts.set(qn("w:eastAsia"), "Calibri")


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run(run, size=16 if level == 1 else 13, bold=True, color=NAVY)
    return p


def set_table_widths(table, widths_cm) -> None:
    """Keep tblGrid and cell tcW in sync so Google Docs / LibreOffice align columns."""
    table.autofit = False
    table.allow_autofit = False
    twips = [int(round(w * TWIPS_PER_CM)) for w in widths_cm]
    total = sum(twips)
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for tw in twips:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(tw))
        grid.append(col)
    tblPr.addnext(grid)
    for row in table.rows:
        for cell, tw, cm in zip(row.cells, twips, widths_cm):
            cell.width = Cm(cm)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(tw))
            tcW.set(qn("w:type"), "dxa")
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn("w:cantSplit")) is None:
            trPr.append(OxmlElement("w:cantSplit"))


def add_code_block(doc, text) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.keep_together = True
    run = p.add_run(text)
    set_run(run, name="Consolas", size=8.5, color=TEXT)
    shade(cell, CODE_HEX)
    set_cell_border(cell, "D0D5DD")
    set_table_widths(table, [CONTENT_WIDTH_CM])
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_callout(doc, text) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    write_cell(cell, text, size=11, bold=True, color=NAVY, fill=AMBER_HEX)
    set_cell_border(cell, AMBER_BORDER)
    set_table_widths(table, [CONTENT_WIDTH_CM])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def add_kv_table(doc, rows, widths=(5.2, 12.2)) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(rows):
        write_cell(table.cell(i, 0), k, bold=True, size=10, fill=ROW_HEX)
        write_cell(table.cell(i, 1), v, size=10)
    set_table_widths(table, list(widths))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_header_table(doc, headers, rows, widths) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        write_cell(table.cell(0, i), h, bold=True, size=9.5, color=WHITE, fill=NAVY_HEX)
    for r, row in enumerate(rows, start=1):
        fill = ROW_HEX if r % 2 == 0 else "FFFFFF"
        for c, value in enumerate(row):
            font = "Consolas" if c in (0, 2, 3) and len(headers) > 3 else "Calibri"
            if len(headers) <= 3 and c == 0:
                font = "Consolas"
            write_cell(table.cell(r, c), value, size=9.5, fill=fill, font=font)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def add_header_footer(doc) -> None:
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = ""
    run = hp.add_run(
        f"SAT Datashare  ·  exp-sat-datashare-prod-api  ·  Kenya  ·  {VERSION} – {VERSION_DATE}"
    )
    set_run(run, size=9, color=NAVY, bold=True)
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.text = ""
    run = fp.add_run(f"Confidential  ·  For the client test team  ·  {VERSION}  ·  Page ")
    set_run(run, size=8.5, color=MUTED)
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run2 = fp.add_run()
    run2._r.append(fld)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run3 = fp.add_run()
    run3._r.append(instr)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run4 = fp.add_run()
    run4._r.append(fld_end)
    for r in (run2, run3, run4):
        set_run(r, size=8.5, color=MUTED)


def token_curl() -> str:
    return (
        f"curl --location '{TOKEN_URL}' \\\n"
        "  --header 'Content-Type: application/x-www-form-urlencoded' \\\n"
        "  --data-urlencode 'grant_type=client_credentials' \\\n"
        f"  --data-urlencode 'client_id={CLIENT_ID}' \\\n"
        "  --data-urlencode 'client_secret=<CLIENT_SECRET>' \\\n"
        f"  --data-urlencode 'scope={SCOPE}'"
    )


def api_curl(path: str, query: str) -> str:
    url = BASE + path
    if query:
        url = f"{url}?{query}"
    return (
        f"curl --location '{url}' \\\n"
        "  --header 'Authorization: Bearer <ACCESS_TOKEN>' \\\n"
        "  --header 'Accept: application/json'"
    )


def sample_payload(name: str) -> str:
    return json.dumps({"data": [SAMPLES[name]]}, indent=2, ensure_ascii=False)


def paged_example() -> str:
    return json.dumps(
        {
            "data": [SAMPLES["Invoice"], "... up to 10,000 rows ..."],
            "next": (
                f"{BASE}/api/v1/Invoice"
                "?fromDate=2026-09-24&toDate=2026-09-24&pageNumber=1"
            ),
        },
        indent=2,
        ensure_ascii=False,
    )


def last_page_example() -> str:
    return json.dumps({"data": [SAMPLES["Invoice"]]}, indent=2, ensure_ascii=False)


def set_core_properties(doc) -> None:
    core = doc.core_properties
    core.author = AUTHOR
    core.last_modified_by = AUTHOR
    core.title = "SAT Datashare Client Test Guide — Kenya"
    core.subject = f"exp-sat-datashare-prod-api {VERSION} – {VERSION_DATE}"
    core.category = "Client test guide"
    core.comments = ""
    core.keywords = ""
    core.revision = 1
    now = datetime.now(timezone.utc)
    core.created = now
    core.modified = now


def strip_generator_metadata(path: Path) -> None:
    """Clear python-docx Application / comments leftovers after save."""
    tmp = path.with_suffix(".tmp.docx")
    with ZipFile(path, "r") as zin, ZipFile(tmp, "w", compression=ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                text = text.replace(">python-docx<", f">{AUTHOR}<")
                text = text.replace("generated by python-docx", "")
                data = text.encode("utf-8")
            elif item.filename == "docProps/app.xml":
                text = data.decode("utf-8")
                text = text.replace(">python-docx<", ">Microsoft Office Word<")
                text = text.replace("generated by python-docx", "")
                data = text.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def build() -> Path:
    doc = Document()
    configure_heading_styles(doc)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    add_header_footer(doc)
    set_core_properties(doc)

    add_para(doc, "CLIENT TEST GUIDE", size=11, bold=True, color=TEAL, space_after=2)
    add_para(doc, "SAT Datashare Experience API", size=22, bold=True, color=NAVY, space_after=2)
    add_para(
        doc,
        f"{VERSION} – {VERSION_DATE}   ·   Kenya   ·   exp-sat-datashare-prod-api",
        size=12,
        color=MUTED,
        space_after=8,
    )
    add_para(
        doc,
        "How to request a Kenya access token and call the eight production GET endpoints.",
        size=11,
        space_after=10,
    )

    add_callout(
        doc,
        "The client secret is not included in this guide. It will be sent separately. "
        "When you receive it, replace <CLIENT_SECRET> in the token request.",
    )

    add_heading(doc, "1. Overview and credentials", level=1)
    add_para(
        doc,
        "Do not send a country, database, or supplier on any request. "
        "The API derives Kenya Reckitt data from your client ID.",
        size=11,
        space_after=8,
    )
    add_kv_table(
        doc,
        [
            ("API name", "exp-sat-datashare-prod-api"),
            ("Base URL", BASE),
            ("Country", "Kenya (derived from your client ID — do not send this)"),
            ("Database", "sat_nobleoutlook (derived from your client ID — do not send this)"),
            ("Supplier", "RECKITT BENCKISER (derived from your client ID — do not send this)"),
            ("Auth", "Azure AD client credentials → Bearer access token"),
            ("Date format", "fromDate / toDate = YYYY-MM-DD   (example: 2026-09-24)"),
            ("Page size", "Up to 10,000 rows per page"),
        ],
    )

    add_para(doc, "Credentials", size=11, bold=True, color=NAVY, space_after=4)
    add_header_table(
        doc,
        ("Field", "Value"),
        [
            ("Tenant ID", TENANT_ID),
            ("Client ID", CLIENT_ID),
            ("Scope", SCOPE),
            ("Client secret", "Sent separately"),
        ],
        [5.2, 12.2],
    )

    add_para(doc, "Placeholders you replace", size=11, bold=True, color=NAVY, space_after=4)
    add_header_table(
        doc,
        ("Placeholder", "Replace with"),
        [
            ("<CLIENT_SECRET>", "Your client secret (sent separately)"),
            ("<ACCESS_TOKEN>", "access_token from the token response"),
        ],
        [5.2, 12.2],
    )

    add_heading(doc, "2. Get an access token", level=1)
    add_para(
        doc,
        "Use Azure AD client credentials. Tokens expire after about 60 minutes "
        "(expires_in: 3599 seconds). Request a new token when calls return HTTP 401.",
        size=11,
        space_after=6,
    )
    for i, step in enumerate(
        (
            "Run the token request below.",
            "Copy access_token from the JSON response.",
            "Send it as the Bearer token on every GET.",
        ),
        start=1,
    ):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(f"{i}.  {step}")
        set_run(run, size=11, color=TEXT)

    add_para(
        doc,
        "Token cURL",
        size=11,
        bold=True,
        color=NAVY,
        space_before=8,
        space_after=4,
        keep_with_next=True,
    )
    add_code_block(doc, token_curl())

    add_para(doc, "Token form fields", size=11, bold=True, color=NAVY, space_after=4)
    add_header_table(
        doc,
        ("Field", "Required", "Value"),
        [
            ("grant_type", "Yes", "client_credentials"),
            ("client_id", "Yes", CLIENT_ID),
            ("client_secret", "Yes", "Your client secret (sent separately)"),
            ("scope", "Yes", SCOPE),
        ],
        [4.4, 2.8, 10.2],
    )

    add_para(
        doc,
        "Successful token response (shape)",
        size=11,
        bold=True,
        color=NAVY,
        space_after=4,
        keep_with_next=True,
    )
    add_code_block(
        doc,
        json.dumps(
            {
                "token_type": "Bearer",
                "expires_in": 3599,
                "ext_expires_in": 3599,
                "access_token": "<ACCESS_TOKEN>",
            },
            indent=2,
        ),
    )

    add_heading(doc, "3. How to call every endpoint", level=1)
    for line in (
        "Use GET. Send only the Authorization header and the query parameters listed for that endpoint.",
        "Authorization: Bearer <ACCESS_TOKEN>",
        "fromDate and toDate are YYYY-MM-DD only. Example: 2026-09-24. Do not send a time (no 00:00:00).",
        "pageNumber is 0-based. First call can omit it or send 0.",
        "A successful call returns HTTP 200 and a JSON object with a data array. Dates inside the payload use dd/MM/yyyy.",
    ):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run("•  " + line)
        set_run(run, size=11, color=TEXT)

    add_heading(doc, "4. Pagination", level=1)
    add_para(
        doc,
        "Each page contains up to 10,000 rows. pageNumber starts at 0. "
        "If another page exists, the body includes a next URL. "
        "The last page has no next field — stop when next is absent.",
        size=11,
        space_after=8,
    )
    add_header_table(
        doc,
        ("Check", "Meaning"),
        [
            ("data array length", "Up to 10,000 rows on this page"),
            ("next is present", "More rows exist. Call the next URL, or increment pageNumber by 1"),
            ("next is absent", "This is the last page. Stop paging"),
            ("pageNumber=0", "First page (you may omit pageNumber)"),
        ],
        [5.2, 12.2],
    )
    add_para(
        doc,
        "Example — more pages remain",
        size=11,
        bold=True,
        color=NAVY,
        space_after=4,
        keep_with_next=True,
    )
    add_code_block(doc, paged_example())
    add_para(
        doc,
        "Example — last page",
        size=11,
        bold=True,
        color=NAVY,
        space_after=4,
        keep_with_next=True,
    )
    add_code_block(doc, last_page_example())
    add_para(
        doc,
        "Follow next with the same Authorization: Bearer <ACCESS_TOKEN> header. "
        "Do not add extra query parameters of your own when using next.",
        size=11,
        space_after=8,
    )

    add_heading(doc, "5. Endpoint catalogue", level=1)
    cat_rows = []
    for i, ep in enumerate(ENDPOINTS, start=1):
        if ep["params"] is DATE_PARAMS:
            qp = "fromDate, toDate, pageNumber"
        elif ep["params"] is PAGE_PARAMS:
            qp = "pageNumber"
        else:
            qp = "None"
        cat_rows.append((str(i), ep["name"], ep["path"], qp))
    add_header_table(
        doc,
        ("#", "Endpoint", "Path", "Query parameters"),
        cat_rows,
        [1.2, 3.4, 5.4, 7.4],
    )

    add_heading(doc, "6. Date format", level=1)
    add_para(
        doc,
        "Query parameters fromDate and toDate must be ISO date-only: YYYY-MM-DD. Both dates are inclusive. "
        "The same calendar day is valid (fromDate=2026-09-24 and toDate=2026-09-24). "
        "Payload dates such as InvoiceDate and OrderDate are returned as dd/MM/yyyy.",
        size=11,
        space_after=6,
    )
    add_header_table(
        doc,
        ("Use", "Format", "Example"),
        [
            ("fromDate / toDate on the URL", "YYYY-MM-DD", "2026-09-24"),
            ("Do not send", "date + time", "2026-09-24T00:00:00  or  24/09/2026"),
            ("Dates inside the JSON body", "dd/MM/yyyy", "24/09/2026"),
        ],
        [6.5, 5.0, 5.9],
    )

    add_heading(doc, "7. Endpoints — cURL and sample output", level=1)
    add_para(
        doc,
        "For each endpoint: copy the cURL, replace <ACCESS_TOKEN>, and confirm HTTP 200. "
        "Sample output is one production row inside the data array.",
        size=11,
        space_after=8,
    )

    for idx, ep in enumerate(ENDPOINTS, start=1):
        add_heading(doc, f"{idx}) {ep['name']}", level=2)
        add_kv_table(
            doc,
            [
                ("Endpoint name", ep["name"]),
                ("Method", "GET"),
                ("URL", BASE + ep["path"]),
                ("Query parameters", "See table below" if ep["params"] else "None"),
            ],
        )
        add_para(doc, ep["notes"], size=10.5, color=MUTED, space_after=6)
        if ep["params"]:
            add_para(doc, "Query parameters", size=11, bold=True, color=NAVY, space_after=4)
            add_header_table(
                doc,
                ("Parameter", "Required", "Format", "Example", "What to provide"),
                list(ep["params"]),
                [2.8, 2.0, 3.2, 2.6, 6.8],
            )
        add_para(doc, "cURL", size=11, bold=True, color=NAVY, space_after=4, keep_with_next=True)
        add_code_block(doc, api_curl(ep["path"], ep["query"]))
        add_para(
            doc,
            "Sample output",
            size=11,
            bold=True,
            color=NAVY,
            space_after=4,
            keep_with_next=True,
        )
        add_code_block(doc, sample_payload(ep["name"]))

    add_heading(doc, "8. Expected results", level=1)
    add_header_table(
        doc,
        ("Check", "Pass when"),
        [
            ("Token call", "HTTP 200 and an access_token string"),
            ("Each GET", "HTTP 200 and a JSON body with a data array"),
            ("Date endpoints", "Rows fall inside the fromDate / toDate window you sent"),
            ("More pages", "next is present and pageNumber can be incremented"),
            ("Last page", "next is absent"),
        ],
        [4.5, 12.9],
    )
    add_para(
        doc,
        "If you need a different date window, change only fromDate and toDate. Keep the YYYY-MM-DD format.",
        size=11,
        space_before=8,
        space_after=4,
    )

    add_heading(doc, "9. Troubleshooting", level=1)
    add_header_table(
        doc,
        ("HTTP status", "Likely cause", "What to do"),
        [
            (
                "401 Unauthorized",
                "Token missing, expired, or from the wrong client",
                "Request a new access token and retry the GET",
            ),
        ],
        [4.2, 6.6, 6.6],
    )

    set_core_properties(doc)
    doc.save(OUT)
    strip_generator_metadata(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
