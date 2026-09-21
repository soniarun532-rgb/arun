#!/usr/bin/env python3
"""Generate SAT datashare Excel mapping workbook and Postman collection/environments."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
MAPPING = ROOT / "mapping"
POSTMAN = ROOT / "postman"

SUPPLIER = "RECKITT BENCKISER"
DATABASE = "sat_nobleoutlook"
DISTRIBUTOR_ID = "1000097205"

INVOICE_SAPI = {
    "ID": 2227,
    "ENTRY_ID": 1640,
    "VISITID": "b4cb0b45-915b-4af6-93b4-fb7b4827f67b",
    "ENTRY_TYPE": "Sale",
    "USERID": 75,
    "SALES_REP": "DICKSON NJUE",
    "REP_CATEGORY": "Van Sales",
    "SUPERVISOR": "SUPERVISOR",
    "SHOPID": 1714,
    "CUSTOMER_ERP_CODE": None,
    "CUSTOMER_NAME": "SYNFORMART",
    "CUSTOMER_ACCOUNT": None,
    "USERCATEGORY": "Van Sales",
    "CUSTOMER_ADDED_DATE": "2026-08-13T16:06:24",
    "CUSTOMER_VERIFICATION": "Valid",
    "CUSTOMER_CATEGORY": "SHOP AND BROWSER",
    "LOCATION_ID": 140,
    "LOCATION_NAME": "Kahawa West",
    "TERRITORY_ID": None,
    "TERRITORY_NAME": None,
    "ROUTE_NAME": "DICKSON NJUE THURSDAY",
    "ROUTE_TYPE": "ON",
    "PRODUCT_ID": 87,
    "REGION_NAME": "NAIROBI",
    "PRODUCT_REF": "BIC9066113",
    "PRODUCT_CODE": "BIC9066113",
    "PRODUCT_CATEGORY": "LIGHTERS",
    "PRODUCT_NAME": "BIC",
    "PRODUCT_SKU": "BIC EZ REACH LIGHTER",
    "PACKAGING_ID": 5,
    "PACKAGING": "CARD",
    "ITEMSWEIGHT": None,
    "SALES_VOLUME": None,
    "QUANTITY": 1,
    "BASE_PRICE": 3563.48,
    "SELLING_PRICE": 3563.48,
    "VALUE_SOLD": 3563.48,
    "SALES_VALUE": 3071.9655172413795,
    "ERP_REFERENCE": None,
    "DELIVERED": "Delivered",
    "ENTRY_TIME": "2026-08-13T16:11:19",
    "TAX_INVOICE_NO": "0111778770000014123",
    "CU_SERIAL_NO": "KRAMW011202209177877",
    "SALE_ORDER_ID": None,
    "STAGE_NAME": None,
    "SUPPLIER_ID": 0,
    "CHANNEL": "Physical Visit",
    "PAYMENT_METHOD": "cash",
    "PAYMENT_STATUS": "Fully Paid",
    "AMOUNT_PAID": 6773.59,
    "LPO_NUMBER": None,
    "PRICELIST_ID": None,
    "PRICE_LIST": None,
    "CREATED_AT": "2026-08-13T16:11:25",
    "UPDATED_AT": "2026-08-13T16:11:25",
    "DISCOUNT": 0,
    "VAT_RATE": 16,
    "PRICE_VAT_INC": 3563.48,
    "PRICE_VAT_EXC": 3071.965517,
    "PRICE_VAT": 491.5144827586,
    "TOTAL_VAT_INC": 3563.48,
    "TOTAL_VAT_EXC": 3071.965517,
    "TOTAL_VAT": 491.5144827586,
    "ROUTEID": 42,
}

INVOICE_PAPI = {
    "STOREID": 1714,
    "Discount": 0,
    "SalesRepID": 75,
    "VATAmount": "491.51",
    "RouteID": 42,
    "DistributorID": DISTRIBUTOR_ID,
    "ProductID": "9066113",
    "WarehouseID": 0,
    "OrderNumber": "",
    "InvoiceNumber": 1640,
    "InvoiceDate": "13/08/2026",
    "InvoiceTime": "16:11:25",
    "InvoiceStatus": "Fully Paid",
    "UnitOfMeasure": "CARD",
    "Currency": "KES",
    "QuantityInvoiced": 1,
    "AmountInvoiced": 3563.48,
    "TransType": "Sale",
}

ORDER_SAPI = {
    "ID": 154,
    "ENTRY_ID": 60,
    "VISIT_ID": "1bca9799-849e-488d-8517-662468133a6f",
    "ENTRY_TYPE": "Order",
    "ERP_REFERENCE": None,
    "USER_ID": 90,
    "SALES_REP": "THOMAS OUMA",
    "REP_CATEGORY": "Van Sales",
    "SUPERVISOR": "SUPERVISOR",
    "SHOP_ID": 1701,
    "CUSTOMER_ERP_CODE": None,
    "CUSTOMER_NAME": "FAVOUR SHOP",
    "CUSTOMER_ACCOUNT": None,
    "USER_CATEGORY": "Van Sales",
    "CUSTOMER_ADDED_DATE": "2026-08-13T13:25:30",
    "CUSTOMER_VERIFICATION": "Valid",
    "CUSTOMER_CATEGORY": "RETAIL SHOP",
    "LOCATION_ID": 92,
    "LOCATION_NAME": "LOWER KABETE",
    "TERRITORY_ID": None,
    "TERRITORY_NAME": None,
    "ROUTE_NAME": "THOMAS OUMA THURSDAY",
    "ROUTE_TYPE": "ON",
    "REGION_NAME": "NAIROBI",
    "PRODUCT_ID": 89,
    "PRODUCT_REF": "BIC976044",
    "PRODUCT_CODE": "BIC976044",
    "PRODUCT_CATEGORY": "SHAVERS",
    "PRODUCT_NAME": "BIC",
    "PRODUCT_SKU": "BIC 1 RAZOR SINGLE POUCH",
    "PACKAGING": "PIECE",
    "QUANTITY": 96,
    "ITEMS_WEIGHT": None,
    "SALES_VOLUME": None,
    "BASE_PRICE": 22.85,
    "SELLING_PRICE": 22.67,
    "VALUE_SOLD": 2176.32,
    "SALES_VALUE": 1876.137931,
    "CHANNEL": "Physical Visit",
    "DELIVERED": "Open",
    "ENTRY_TIME": "2026-08-20T14:03:22",
    "STAGE_NAME": "New Order",
    "SUPPLIER_ID": 183,
    "LPO_NUMBER": None,
    "PRICELIST_ID": None,
    "PRICE_LIST": None,
    "CREATED_AT": "2026-08-20T14:03:24",
    "UPDATED_AT": "2026-08-20T14:03:24",
    "DISCOUNT": 17.28,
    "VAT_RATE": 16,
    "PRICE_VAT_INC": 22.67,
    "PRICE_VAT_EXC": 19.543103,
    "PRICE_VAT": 3.1268965517,
    "TOTAL_VAT_INC": 2176.32,
    "TOTAL_VAT_EXC": 1876.137931,
    "TOTAL_VAT": 300.1820689655,
    "ROUTEID": 85,
    "CHECKINTIME": "2026-08-20T14:02:14",
    "CHECKOUTTIME": "2026-08-20T14:03:24",
}

ORDER_PAPI = {
    "StoreID": 1701,
    "SalesRepID": 90,
    "RouteID": 85,
    "ProductID": "976044",
    "DistributorID": "???",
    "WarehouseID": 183,
    "OrderNumber": 60,
    "OrderDate": "20/08/2026",
    "OrderStartTime": "14:02:14",
    "OrderEndTime": "14:03:24",
    "OrderStatus": "Open",
    "UnitOfMeasure": "PIECE",
    "Currency": "KES",
    "QuantityOrdered": 96,
    "AmountOrdered": 2176.32,
    "Discount": 17.28,
    "VATAmount": "300.18",
}

PRODUCT_SAPI = {
    "id": 1,
    "productcode": "BIC502418",
    "product_category": "STATIONERY",
    "product_name": "BIC",
    "product_desc": "CRISTAL MED BLU BX 25 BCL KE",
    "product_status": "Enabled",
    "short_code": "BIC502418",
    "tax_code": "A",
    "hs_code": None,
    "focus_product": 0,
    "tonne_equivalent": None,
    "apply_discount": None,
    "disable_price_check": False,
    "product_type": "item",
    "supplier": SUPPLIER,
    "alternative_group": None,
    "supplier_id": 2,
    "uomname": "PACKET",
}

PRODUCT_PAPI = {
    "ProductID": "502418",
    "ProductSKU": "CRISTAL MED BLU BX 25 BCL KE",
    "UnitofMeasure": "PACKET",
}

ROUTE_SAPI = {
    "id": 31,
    "route_id": 31,
    "region_id": 1,
    "region_name": "NAIROBI",
    "route_name": "DICKSON NJUE MONDAY",
    "routestatus": "Active",
    "created_by": "CLEMENCEA SARU",
    "updated_by": None,
    "created_at": "2026-06-15T11:15:47",
    "updated_at": "2026-06-15T11:15:47",
    "userid": 75,
    "visit_frequency": "Weekly",
    "visit_week": "1",
    "visit_day": "MON",
    "status": "Active",
    "saler_category": "Van Sales",
}

ROUTE_PAPI = {
    "RouteID": 31,
    "RouteName": "DICKSON NJUE MONDAY",
    "SalesRepID": 75,
    "RouteType": "Weekly",
    "SalesRepType": "Van Sales",
    "Date": "15/06/2026",
}

SALESREP_SAPI = {
    "id": 72,
    "user_reference": None,
    "warehouse_code": None,
    "user_type": 2,
    "user_channel": 1,
    "region_id": 1,
    "rep_category": 4,
    "name": "ESTHER KAGIRI",
    "email": "ngong@nobleoutlook.co.ke",
    "phone_number": "0710000000",
    "last_time_active": "3",
    "trackactivities": "YES",
    "demoaccount": "",
    "billing": "billable",
    "status": "Inactive",
    "territory_id": None,
    "rolegroup_id": 4,
    "vehicle_id": 4,
    "added_by": 3,
    "created_at": "2026-06-12T06:40:33",
    "saler_category": "Drivers (Delivery Agents)",
    "STOCKPOINT_ID": 183,
}

SALESREP_PAPI = {
    "SalesRepID": 72,
    "SalesRepName": "ESTHER KAGIRI",
    "DistributorID": DISTRIBUTOR_ID,
    "WarehouseID": 183,
    "SalesRepType": "Drivers (Delivery Agents)",
}

CUSTOMER_SAPI = {
    "ID": 199,
    "SHOPID": 199,
    "PARENT_ID": None,
    "region_id": 1,
    "group_id": 0,
    "supplied_by": 183,
    "shop_cat_id": 2,
    "shop_subcat_id": 2,
    "CUSTOMERCODE": None,
    "CUSTOMERNAME": "MANDELA SHOP",
    "IS_HQ": "no",
    "ACCOUNT_ID": 0,
    "ACCOUNT": None,
    "REGIONNAME": "NAIROBI",
    "LOCATION_ID": 85,
    "LOCATIONNAME": "RUIRU",
    "CUSTOMERCATEGORY": "RETAIL SHOP",
    "TERRITORY_NAME": None,
    "CHANNEL": "GENERAL TRADE",
    "RELATIONSHIP": "INDIRECT",
    "SUPPLIER": "GIKAMBURA STOCKPOINT",
    "PHONENUMBER": "+254722590524",
    "VERIFIED": "Valid",
    "VERIFIEDBY": "IAN MWENDWA",
    "VERIFIEDDATE": "2026-06-15T14:35:26",
    "ADDEDBY": "PHILLIP DAVID",
    "USERID": 73,
    "USERCATEGORY": "Van Sales",
    "EMAILADDRESS": "",
    "CONTACTPERSON": "PINCH ANISO",
    "KRAPIN": None,
    "LASTVISIT": "2026-09-07T15:27:43",
    "LASTSALE": "2026-06-15T15:07:50",
    "LATITUDE": -1.2349905,
    "LONGITUDE": 36.9328219,
    "DATECREATED": "2026-06-15T14:17:10",
    "STATUS": 1,
    "DISTRIBUTOR_ID": None,
    "COUNTRY_ID": 113,
    "COUNTRY_NAME": "KENYA",
    "CONVERTED_LEAD": "Yes",
    "CREDIT_LIMIT": 0,
    "PAYMENT_TERMS": None,
    "CREDIT_DAYS": None,
    "CUSTOMER_TYPE": "Customer",
}

CUSTOMER_PAPI = {
    "StoreID": 199,
    "StoreName": "MANDELA SHOP",
    "StoreCategory": "RETAIL SHOP",
    "Channel": "GENERAL TRADE",
    "StoreClassification": "GENERAL TRADE",
    "WarehouseID": "???",
    "Longitude": 36.9328219,
    "Latitude": -1.2349905,
    "City": "RUIRU",
    "Region": "NAIROBI",
    "County": "",
    "StoreStatus": 1,
    "StoreCreationDate": "15/06/2026",
    "StoreSize": "",
}

STOCK_SAPI = {
    "productcode": "BIC969584",
    "uomname": "PACKET",
    "stockpoint_name": None,
    "date_created": "2026-08-17T14:00:56",
    "quantity": 2640,
    "total_cost": 1060,
}

STOCK_PAPI = {
    "ProductSKU": "969584",
    "UnitofMeasure": "PACKET",
    "DistributorWarehouseCode": "",
    "Date": "17/08/2026",
    "InventoryQuantity": 2640,
    "InventoryPrice": 1060,
}

WAREHOUSE_SAPI = {
    "ID": 183,
    "SHOPID": 183,
    "PARENT_ID": None,
    "region_id": 1,
    "group_id": 0,
    "supplied_by": 0,
    "shop_cat_id": 1,
    "shop_subcat_id": None,
    "CUSTOMERCODE": "183",
    "CUSTOMERNAME": "GIKAMBURA STOCKPOINT",
    "IS_HQ": "no",
    "ACCOUNT_ID": 0,
    "ACCOUNT": None,
    "REGIONNAME": "NAIROBI",
    "LOCATION_ID": 4,
    "LOCATIONNAME": "GIKAMBURA",
    "CUSTOMERCATEGORY": "WHOLESALE",
    "TERRITORY_NAME": None,
    "CHANNEL": "GENERAL TRADE",
    "RELATIONSHIP": "DIRECT",
    "SUPPLIER": None,
    "PHONENUMBER": None,
    "VERIFIED": "Pending",
    "VERIFIEDBY": None,
    "VERIFIEDDATE": None,
    "ADDEDBY": "TECLA AYORE",
    "USERID": 3,
    "USERCATEGORY": "Admin",
    "EMAILADDRESS": "",
    "CONTACTPERSON": None,
    "KRAPIN": None,
    "LASTVISIT": None,
    "LASTSALE": None,
    "LATITUDE": None,
    "LONGITUDE": None,
    "DATECREATED": "2026-06-08T15:07:59",
    "STATUS": 1,
    "DISTRIBUTOR_ID": None,
    "COUNTRY_ID": 113,
    "COUNTRY_NAME": "KENYA",
    "CONVERTED_LEAD": "Yes",
    "CREDIT_LIMIT": 0,
    "PAYMENT_TERMS": None,
    "CREDIT_DAYS": None,
    "CUSTOMER_TYPE": "Supplier",
}

WAREHOUSE_PAPI = {
    "WarehouseID": 183,
    "WarehouseName": "GIKAMBURA STOCKPOINT",
    "DistributorID": "",
}

SQL_PAGE = "LIMIT <pageSize + 1> OFFSET <pageNumber * pageSize>"

INVOICE_SQL = f"""Select a.*, b.routeid
from {{database}}.bi_salesmaster as a
left join {{database}}.bi_customer_visits as b on a.VisitID = b.visitid
inner join {{database}}.bi_products as c on a.PRODUCT_ID = c.id and c.supplier = '{{supplier}}'
where a.Created_AT between '<fromDate> 00:00:00' and '<toDate> 00:00:00'
order by a.Entry_ID desc
{SQL_PAGE}"""

ORDER_SQL = f"""Select a.*, b.routeid, b.CHECKINTIME, b.CHECKOUTTIME
from {{database}}.bi_ordersmaster as a
inner join {{database}}.bi_products as c on a.PRODUCT_ID = c.id and c.supplier = '{{supplier}}'
left join {{database}}.bi_customer_visits as b on a.VISIT_ID = b.visitid
order by a.ENTRY_ID desc
{SQL_PAGE}

Note: Exp/PAPI still send fromDate and toDate. Current SAPI Order SQL does not filter on those dates."""

PRODUCT_SQL = f"""Select a.*, c.uomname
from {{database}}.bi_products as a
left join {{database}}.bi_uom_quantities as b on a.id = b.product_id
left join {{database}}.bi_uom as c on b.packaging_id = c.id
where a.supplier = '{{supplier}}'
order by a.id
{SQL_PAGE}"""

ROUTE_SQL = f"""Select a.*, b.userid, b.visit_frequency, b.visit_week, b.visit_day, b.status, d.saler_category
from {{database}}.bi_routes as a
left join {{database}}.bi_user_routes as b on a.route_id = b.route_id
left join {{database}}.bi_users as c on b.userid = c.id
left join {{database}}.bi_user_categories as d on c.rep_category = d.usercategory
where exists (
  Select 1 from (
    Select distinct routeid from (
      select Distinct v.routeid, y.product_id
      from {{database}}.bi_salesmaster as y
      inner join {{database}}.bi_customer_visits as v on y.VisitID = v.visitid
      union all
      select Distinct v.routeid, y.product_id
      from {{database}}.bi_ordersmaster as y
      inner join {{database}}.bi_customer_visits as v on y.Visit_ID = v.visitid
    ) as x
    inner join {{database}}.bi_products as z on x.product_id = z.id
    where z.supplier = '{{supplier}}'
  ) as w
  where w.routeid = a.route_id
)
order by a.route_id
{SQL_PAGE}"""

SALESREP_SQL = f"""Select a.*, b.saler_category, c.STOCKPOINT_ID
from {{database}}.bi_users as a
left join {{database}}.bi_user_categories as b on a.rep_category = b.usercategory
left join (
  Select bia.REP_ID, bia.STOCKPOINT_ID
  from {{database}}.bi_stockpoint_allocation bia
  inner join {{database}}.bi_customer_master bicm on bia.stockpoint_ID = bicm.SHOPID
  where CUSTOMER_TYPE = 'Supplier'
    and (UPPER(CustomerName) like '%STOCKPOINT' or UPPER(CustomerName) like '%WAREHOUSE')
) as c on a.id = c.rep_id
where exists (
  Select 1 from (
    Select distinct userid from (
      select Distinct userid, y.product_id from {{database}}.bi_salesmaster as y
      union all
      select Distinct user_id as userid, y.product_id from {{database}}.bi_ordersmaster as y
    ) as x
    inner join {{database}}.bi_products as z on x.product_id = z.id
    where z.supplier = '{{supplier}}'
  ) as w
  where w.userid = a.id
)
order by a.id
{SQL_PAGE}

Change vs previous: join bi_stockpoint_allocation + supplier customers to return STOCKPOINT_ID."""

CUSTOMER_SQL = f"""Select a.*
from {{database}}.bi_customer_master as a
where a.Customer_Type = 'Customer'
and exists (
  Select 1 from (
    Select distinct SHOPID from (
      select Distinct SHOPID, product_id from {{database}}.bi_salesmaster as y
      union all
      select Distinct SHOP_ID as SHOPID, product_id from {{database}}.bi_ordersmaster as y
    ) as x
    inner join {{database}}.bi_products as z on x.product_id = z.id
    where z.supplier = '{{supplier}}'
  ) as w
  where w.SHOPID = a.SHOPID
)
order by a.SHOPID
{SQL_PAGE}

Note: Exp/PAPI send fromDate and toDate. Current SAPI Customer SQL does not filter on those dates."""

STOCK_SQL = f"""Select a.productcode, d.UOMName, b.stockpoint_name, MAX(b.date_created) as date_created,
SUM(CASE WHEN LOWER(b.entry_type) = 'goodsin' THEN b.quantity ELSE 0 END)
  - SUM(CASE WHEN LOWER(b.entry_type) = 'goodsout' THEN b.quantity ELSE 0 END) as quantity,
SUM(b.total_cost) as total_cost
from {{database}}.bi_products as a
inner join {{database}}.bi_stock_movement as b on a.id = b.product_id
left join {{database}}.bi_uom_quantities as c on b.product_id = c.product_id
left join {{database}}.bi_uom as d on c.packaging_id = d.id
where a.supplier = '{{supplier}}'
group by a.id, d.UOMName, b.stockpoint_name
order by a.id, d.UOMName, b.stockpoint_name
{SQL_PAGE}

Change vs previous: quantity = goodsin minus goodsout. Do not SUM every movement row."""

WAREHOUSE_SQL = f"""Select c.*
from {{database}}.bi_customer_master as c
where c.CUSTOMER_TYPE = 'Supplier'
/* and exists (Select 1 from (select distinct stockpoint_id from {{database}}.bi_products as a
inner join {{database}}.bi_stock_movement as b on a.id = b.product_id
where a.supplier = '{{supplier}}') as z where z.stockpoint_id = c.id) */
order by c.SHOPID
{SQL_PAGE}

EXISTS on stockpoint_id stays commented until Solutech adds the column."""

INVOICE_MAP = """SHOPID -> STOREID
DISCOUNT -> Discount
USERID -> SalesRepID
TOTAL_VAT -> VATAmount  (PAPI amount2, 2 decimals, format 0.00)
ROUTEID -> RouteID
papi.sat.solutech.distributor.id (1000097205) -> DistributorID
PRODUCT_CODE -> ProductID  (strip leading letters)
SUPPLIER_ID -> WarehouseID
SALE_ORDER_ID -> OrderNumber
ENTRY_ID -> InvoiceNumber
CREATED_AT -> InvoiceDate  (dd/MM/yyyy)
CREATED_AT -> InvoiceTime  (HH:mm:ss)
PAYMENT_STATUS -> InvoiceStatus
PACKAGING -> UnitOfMeasure
papi.sat.solutech.currency (KES) -> Currency
QUANTITY -> QuantityInvoiced
VALUE_SOLD -> AmountInvoiced
ENTRY_TYPE -> TransType"""

ORDER_MAP = """SHOP_ID -> StoreID
USER_ID -> SalesRepID  (not SALES_REP name)
ROUTEID -> RouteID
PRODUCT_CODE -> ProductID  (strip leading letters)
hardcoded '???' -> DistributorID
SUPPLIER_ID -> WarehouseID
ENTRY_ID -> OrderNumber
CHECKINTIME -> OrderDate  (dd/MM/yyyy)
CHECKINTIME -> OrderStartTime  (HH:mm:ss)
CHECKOUTTIME -> OrderEndTime  (HH:mm:ss)
DELIVERED -> OrderStatus
PACKAGING -> UnitOfMeasure
papi.sat.solutech.currency (KES) -> Currency
QUANTITY -> QuantityOrdered
TOTAL_VAT_INC -> AmountOrdered
DISCOUNT -> Discount
TOTAL_VAT -> VATAmount  (PAPI amount2, 2 decimals)"""

PRODUCT_MAP = """productcode -> ProductID  (strip leading letters)
product_desc -> ProductSKU
uomname -> UnitofMeasure

PAPI returns only these three fields."""

ROUTE_MAP = """route_id -> RouteID
route_name -> RouteName
userid -> SalesRepID
visit_frequency -> RouteType
saler_category -> SalesRepType
created_at -> Date  (dd/MM/yyyy)"""

SALESREP_MAP = """id -> SalesRepID
name -> SalesRepName
papi.sat.solutech.distributor.id (1000097205) -> DistributorID
STOCKPOINT_ID -> WarehouseID  (from bi_stockpoint_allocation join; not warehouse_code)
saler_category -> SalesRepType"""

CUSTOMER_MAP = """SHOPID -> StoreID
CUSTOMERNAME -> StoreName
CUSTOMERCATEGORY -> StoreCategory
CHANNEL -> Channel
CHANNEL -> StoreClassification
hardcoded '???' -> WarehouseID
LONGITUDE -> Longitude
LATITUDE -> Latitude
LOCATIONNAME -> City
REGIONNAME -> Region
'' -> County
STATUS -> StoreStatus
DATECREATED -> StoreCreationDate  (dd/MM/yyyy)
'' -> StoreSize"""

STOCK_MAP = """productcode -> ProductSKU  (strip leading letters)
uomname -> UnitofMeasure
stockpoint_name -> DistributorWarehouseCode
date_created -> Date  (dd/MM/yyyy)
quantity -> InventoryQuantity  (goodsin minus goodsout)
total_cost -> InventoryPrice"""

WAREHOUSE_MAP = """ID -> WarehouseID
CUSTOMERNAME -> WarehouseName
'' -> DistributorID"""

RESOURCES = [
    {
        "name": "Invoice",
        "sapi_path": "/api/invoice",
        "papi_path": "/api/Adapter_Invoice",
        "exp_path": "/api/v1/Invoice",
        "dwl": "dwl/map-adapter-invoice.dwl",
        "sapi": INVOICE_SAPI,
        "papi": INVOICE_PAPI,
        "mapping": INVOICE_MAP,
        "sql": INVOICE_SQL,
        "status": "UPDATED — DistributorID 1000097205, VAT 2 decimals, ProductID numeric",
        "exp_params": "fromDate, toDate, pageNumber",
        "papi_params": "database, supplier, fromDate, toDate, pageNumber",
        "sapi_params": "database, supplier, fromDate, toDate, pageNumber",
        "date_used_in_sql": "Yes — Created_AT between fromDate 00:00:00 and toDate 00:00:00",
    },
    {
        "name": "Order",
        "sapi_path": "/api/order",
        "papi_path": "/api/Adapter_Order",
        "exp_path": "/api/v1/Order",
        "dwl": "dwl/map-adapter-order.dwl",
        "sapi": ORDER_SAPI,
        "papi": ORDER_PAPI,
        "mapping": ORDER_MAP,
        "sql": ORDER_SQL,
        "status": "UPDATED — SalesRepID = USER_ID, ProductID numeric, VAT 2 decimals",
        "exp_params": "fromDate, toDate, pageNumber",
        "papi_params": "database, supplier, fromDate, toDate, pageNumber",
        "sapi_params": "database, supplier, fromDate, toDate, pageNumber",
        "date_used_in_sql": "No — params accepted, SQL does not filter dates",
    },
    {
        "name": "Product",
        "sapi_path": "/api/product",
        "papi_path": "/api/Adapter_Product",
        "exp_path": "/api/v1/Product",
        "dwl": "dwl/map-adapter-product.dwl",
        "sapi": PRODUCT_SAPI,
        "papi": PRODUCT_PAPI,
        "mapping": PRODUCT_MAP,
        "sql": PRODUCT_SQL,
        "status": "UPDATED — only ProductID, ProductSKU, UnitofMeasure",
        "exp_params": "(none)",
        "papi_params": "database, supplier, pageNumber",
        "sapi_params": "database, supplier, pageNumber",
        "date_used_in_sql": "N/A — no date params",
    },
    {
        "name": "Route",
        "sapi_path": "/api/route",
        "papi_path": "/api/Adapter_Route",
        "exp_path": "/api/v1/Route",
        "dwl": "dwl/map-adapter-route.dwl",
        "sapi": ROUTE_SAPI,
        "papi": ROUTE_PAPI,
        "mapping": ROUTE_MAP,
        "sql": ROUTE_SQL,
        "status": "Unchanged mapping",
        "exp_params": "(none)",
        "papi_params": "database, supplier, pageNumber",
        "sapi_params": "database, supplier, pageNumber",
        "date_used_in_sql": "N/A — no date params",
    },
    {
        "name": "SalesRep",
        "sapi_path": "/api/sales-rep",
        "papi_path": "/api/Adapter_SalesRep",
        "exp_path": "/api/v1/SalesRep",
        "dwl": "dwl/map-adapter-salesrep.dwl",
        "sapi": SALESREP_SAPI,
        "papi": SALESREP_PAPI,
        "mapping": SALESREP_MAP,
        "sql": SALESREP_SQL,
        "status": "UPDATED — STOCKPOINT_ID join, DistributorID 1000097205",
        "exp_params": "(none)",
        "papi_params": "database, supplier, pageNumber",
        "sapi_params": "database, supplier, pageNumber",
        "date_used_in_sql": "N/A — no date params",
    },
    {
        "name": "Customer",
        "sapi_path": "/api/customer",
        "papi_path": "/api/Adapter_Customer",
        "exp_path": "/api/v1/Customer",
        "dwl": "dwl/map-adapter-customer.dwl",
        "sapi": CUSTOMER_SAPI,
        "papi": CUSTOMER_PAPI,
        "mapping": CUSTOMER_MAP,
        "sql": CUSTOMER_SQL,
        "status": "Unchanged mapping",
        "exp_params": "fromDate, toDate, pageNumber",
        "papi_params": "database, supplier, fromDate, toDate, pageNumber",
        "sapi_params": "database, supplier, fromDate, toDate, pageNumber",
        "date_used_in_sql": "No — params accepted, SQL does not filter dates",
    },
    {
        "name": "Stock",
        "sapi_path": "/api/stock",
        "papi_path": "/api/Adapter_Stock",
        "exp_path": "/api/v1/Stock",
        "dwl": "dwl/map-adapter-stock.dwl",
        "sapi": STOCK_SAPI,
        "papi": STOCK_PAPI,
        "mapping": STOCK_MAP,
        "sql": STOCK_SQL,
        "status": "UPDATED — quantity = goodsin minus goodsout",
        "exp_params": "pageNumber",
        "papi_params": "database, supplier, pageNumber",
        "sapi_params": "database, supplier, pageNumber",
        "date_used_in_sql": "N/A — no date params",
    },
    {
        "name": "Warehouse",
        "sapi_path": "/api/warehouse",
        "papi_path": "/api/Adapter_Warehouse",
        "exp_path": "/api/v1/Warehouse",
        "dwl": "dwl/map-adapter-warehouse.dwl",
        "sapi": WAREHOUSE_SAPI,
        "papi": WAREHOUSE_PAPI,
        "mapping": WAREHOUSE_MAP,
        "sql": WAREHOUSE_SQL,
        "status": "Unchanged mapping. EXISTS still commented.",
        "exp_params": "(none)",
        "papi_params": "database, supplier, pageNumber",
        "sapi_params": "database, supplier, pageNumber",
        "date_used_in_sql": "N/A — no date params",
    },
]


def dumps(obj) -> str:
    return json.dumps(obj, indent=2, ensure_ascii=False)


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
CELL_FONT = Font(name="Calibri", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
ALT_FILL = PatternFill("solid", fgColor="F2F2F2")


def style_header(ws, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_rows(ws, headers, rows, widths):
    ws.append(headers)
    for i, row in enumerate(rows, start=2):
        ws.append(row)
        for col in range(1, len(headers) + 1):
            cell = ws.cell(i, col)
            cell.font = CELL_FONT
            cell.alignment = WRAP
            cell.border = THIN
            if i % 2 == 0:
                cell.fill = ALT_FILL
    style_header(ws, len(headers))
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 22
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 90 if ws.max_column > 4 else 30


def build_excel(path: Path):
    wb = Workbook()

    payloads = wb.active
    payloads.title = "Payloads"
    write_rows(
        payloads,
        [
            "Resource",
            "SAPI payload",
            "Transformed payload",
            "Mapping",
            "Query",
            "SAPI endpoint",
            "PAPI endpoint",
            "Exp endpoint",
            "PAPI mapping file",
            "Status",
        ],
        [
            [
                r["name"],
                dumps(r["sapi"]),
                dumps(r["papi"]),
                r["mapping"],
                r["sql"],
                r["sapi_path"],
                r["papi_path"],
                r["exp_path"],
                r["dwl"],
                r["status"],
            ]
            for r in RESOURCES
        ],
        [14, 48, 42, 48, 58, 18, 24, 18, 32, 42],
    )
    for r in range(2, payloads.max_row + 1):
        payloads.row_dimensions[r].height = 220

    endpoints = wb.create_sheet("Endpoints")
    write_rows(
        endpoints,
        [
            "Resource",
            "SAPI endpoint",
            "PAPI endpoint",
            "Exp (EAPI) endpoint",
            "PAPI mapping file",
            "Mapping status",
            "Exp query params",
            "PAPI / SAPI query params",
            "Dates used in SQL?",
        ],
        [
            [
                r["name"],
                r["sapi_path"],
                r["papi_path"],
                r["exp_path"],
                r["dwl"],
                r["status"],
                r["exp_params"],
                r["papi_params"],
                r["date_used_in_sql"],
            ]
            for r in RESOURCES
        ],
        [14, 18, 24, 22, 32, 48, 32, 48, 48],
    )
    for r in range(2, endpoints.max_row + 1):
        endpoints.row_dimensions[r].height = 48

    field_rows = []
    field_specs = {
        "Invoice": [
            ("STOREID", "SHOPID", "direct"),
            ("Discount", "DISCOUNT", "direct"),
            ("SalesRepID", "USERID", "direct"),
            ("VATAmount", "TOTAL_VAT", "amount2 format 0.00"),
            ("RouteID", "ROUTEID", "direct"),
            ("DistributorID", "papi.sat.solutech.distributor.id", "hardcoded 1000097205"),
            ("ProductID", "PRODUCT_CODE", "strip leading letters"),
            ("WarehouseID", "SUPPLIER_ID", "direct"),
            ("OrderNumber", "SALE_ORDER_ID", "null becomes empty"),
            ("InvoiceNumber", "ENTRY_ID", "direct"),
            ("InvoiceDate", "CREATED_AT", "dd/MM/yyyy"),
            ("InvoiceTime", "CREATED_AT", "HH:mm:ss"),
            ("InvoiceStatus", "PAYMENT_STATUS", "direct"),
            ("UnitOfMeasure", "PACKAGING", "direct"),
            ("Currency", "papi.sat.solutech.currency", "hardcoded KES"),
            ("QuantityInvoiced", "QUANTITY", "direct"),
            ("AmountInvoiced", "VALUE_SOLD", "direct"),
            ("TransType", "ENTRY_TYPE", "direct"),
        ],
        "Order": [
            ("StoreID", "SHOP_ID", "direct"),
            ("SalesRepID", "USER_ID", "not SALES_REP"),
            ("RouteID", "ROUTEID", "direct"),
            ("ProductID", "PRODUCT_CODE", "strip leading letters"),
            ("DistributorID", "hardcoded", "'???' still in PAPI DWL"),
            ("WarehouseID", "SUPPLIER_ID", "direct"),
            ("OrderNumber", "ENTRY_ID", "direct"),
            ("OrderDate", "CHECKINTIME", "dd/MM/yyyy"),
            ("OrderStartTime", "CHECKINTIME", "HH:mm:ss"),
            ("OrderEndTime", "CHECKOUTTIME", "HH:mm:ss"),
            ("OrderStatus", "DELIVERED", "direct"),
            ("UnitOfMeasure", "PACKAGING", "direct"),
            ("Currency", "papi.sat.solutech.currency", "hardcoded KES"),
            ("QuantityOrdered", "QUANTITY", "direct"),
            ("AmountOrdered", "TOTAL_VAT_INC", "direct"),
            ("Discount", "DISCOUNT", "direct"),
            ("VATAmount", "TOTAL_VAT", "amount2 format 0.00"),
        ],
        "Product": [
            ("ProductID", "productcode", "strip leading letters"),
            ("ProductSKU", "product_desc", "direct"),
            ("UnitofMeasure", "uomname", "direct"),
        ],
        "Route": [
            ("RouteID", "route_id", "direct"),
            ("RouteName", "route_name", "direct"),
            ("SalesRepID", "userid", "direct"),
            ("RouteType", "visit_frequency", "direct"),
            ("SalesRepType", "saler_category", "direct"),
            ("Date", "created_at", "dd/MM/yyyy"),
        ],
        "SalesRep": [
            ("SalesRepID", "id", "direct"),
            ("SalesRepName", "name", "direct"),
            ("DistributorID", "papi.sat.solutech.distributor.id", "hardcoded 1000097205"),
            ("WarehouseID", "STOCKPOINT_ID", "from bi_stockpoint_allocation join"),
            ("SalesRepType", "saler_category", "direct"),
        ],
        "Customer": [
            ("StoreID", "SHOPID", "direct"),
            ("StoreName", "CUSTOMERNAME", "direct"),
            ("StoreCategory", "CUSTOMERCATEGORY", "direct"),
            ("Channel", "CHANNEL", "direct"),
            ("StoreClassification", "CHANNEL", "same as Channel"),
            ("WarehouseID", "hardcoded", "'???' still in PAPI DWL"),
            ("Longitude", "LONGITUDE", "direct"),
            ("Latitude", "LATITUDE", "direct"),
            ("City", "LOCATIONNAME", "direct"),
            ("Region", "REGIONNAME", "direct"),
            ("County", "", "empty"),
            ("StoreStatus", "STATUS", "direct"),
            ("StoreCreationDate", "DATECREATED", "dd/MM/yyyy"),
            ("StoreSize", "", "empty"),
        ],
        "Stock": [
            ("ProductSKU", "productcode", "strip leading letters"),
            ("UnitofMeasure", "uomname", "direct"),
            ("DistributorWarehouseCode", "stockpoint_name", "direct"),
            ("Date", "date_created", "dd/MM/yyyy"),
            ("InventoryQuantity", "quantity", "goodsin minus goodsout"),
            ("InventoryPrice", "total_cost", "direct"),
        ],
        "Warehouse": [
            ("WarehouseID", "ID", "direct"),
            ("WarehouseName", "CUSTOMERNAME", "direct"),
            ("DistributorID", "", "empty"),
        ],
    }
    for resource, fields in field_specs.items():
        for dest, src, transform in fields:
            field_rows.append([resource, dest, src, transform])

    field_map = wb.create_sheet("Current field map")
    write_rows(
        field_map,
        ["Resource", "Final PAPI / Exp field", "SAPI source column", "Transform"],
        field_rows,
        [16, 28, 40, 40],
    )
    for r in range(2, field_map.max_row + 1):
        field_map.row_dimensions[r].height = 22

    query_rows = []
    for r in RESOURCES:
        query_rows.append(
            [
                r["name"],
                "Exp",
                r["exp_path"],
                r["exp_params"],
                "Authorization: Bearer <Azure access_token>",
                "Clients never send database or supplier. Exp injects them from the JWT appid map.",
            ]
        )
        query_rows.append(
            [
                r["name"],
                "PAPI",
                r["papi_path"],
                r["papi_params"],
                "Authorization: Test  (hardcoded any value, not a variable)",
                "database and supplier required. pageNumber is a string when forwarded to SAPI.",
            ]
        )
        query_rows.append(
            [
                r["name"],
                "SAPI",
                r["sapi_path"],
                r["sapi_params"],
                "No Authorization header",
                "database = MySQL schema. supplier = products.supplier. pageNumber 0-based. page size 10000.",
            ]
        )

    query_ws = wb.create_sheet("Query parameters")
    write_rows(
        query_ws,
        ["Resource", "Layer", "Path", "Query parameters", "Headers", "Notes"],
        query_rows,
        [14, 10, 26, 52, 52, 70],
    )
    for r in range(2, query_ws.max_row + 1):
        query_ws.row_dimensions[r].height = 36

    auth_rows = [
        [
            "Exp",
            "CloudHub / API Manager",
            "OpenID Connect Access Token Enforcement",
            "Authorization: Bearer {{access_token}}",
            "Azure AD client_credentials. Token URL https://login.microsoftonline.com/f4faf003-d90b-4832-9df8-c8a22d29bdd4/oauth2/v2.0/token. Scope https://graph.microsoft.com/.default. Kenya client id de772600-0d1f-4492-90b0-a57fcc284cf9. Autodiscovery api.id=2850714 on flow exp-sat-datashare-api-main (exp-sat-datashare-sand-api).",
        ],
        [
            "Exp",
            "Localhost Studio",
            "Same Bearer token. Policy may not run locally.",
            "Authorization: Bearer {{access_token}}",
            "Exp still decodes JWT appid/azp and maps to exp.sat.oidc.<country>.database and .supplier. Secret is not stored in the API.",
        ],
        [
            "Exp tenant map",
            "All",
            "JWT appid -> country",
            "Do not send client_id, client_secret, database, or supplier on the GET",
            "kenya -> sat_nobleoutlook / RECKITT BENCKISER. libya sat_salvogrima, malta sat_sgmalta, uganda sat_sgug (client ids TBD).",
        ],
        [
            "PAPI",
            "All environments",
            "Header present. Any value accepted by the current policy.",
            "Authorization: Test",
            "Hardcoded. No Postman variable. Exp already sends Authorization: Test when it calls PAPI.",
        ],
        [
            "SAPI",
            "All environments",
            "None for these eight GETs",
            "(no Authorization header)",
            "PAPI internally sends Bearer test to SAPI. Direct Postman SAPI calls do not need a header.",
        ],
    ]
    auth_ws = wb.create_sheet("Authentication")
    write_rows(
        auth_ws,
        ["Layer", "Where", "Mechanism", "Header to send", "Detail"],
        auth_rows,
        [16, 24, 44, 44, 80],
    )
    for r in range(2, auth_ws.max_row + 1):
        auth_ws.row_dimensions[r].height = 70

    env_ws = wb.create_sheet("Environments")
    write_rows(
        env_ws,
        ["Environment", "exp_base", "papi_base", "sapi_base", "Notes"],
        [
            [
                "local",
                "http://localhost:8081",
                "http://localhost:8081",
                "http://localhost:8081",
                "Run one API at a time in Studio. Switch the matching folder.",
            ],
            [
                "sandbox",
                "https://exp-sat-datashare-sand-api-5nct48.2ky31l-2.deu-c1.eu1.cloudhub.io",
                "https://prc-solutech-api-v1-5nct48.2ky31l-1.deu-c1.eu1.cloudhub.io",
                "https://sys-solutech-api-v1-5nct48.2ky31l-2.deu-c1.eu1.cloudhub.io",
                "Current CloudHub sandbox apps.",
            ],
            [
                "presandbox",
                "",
                "",
                "",
                "Paste hosts when the presandbox apps are deployed. Collection uses {{exp_base}} / {{papi_base}} / {{sapi_base}}.",
            ],
            [
                "production",
                "",
                "",
                "",
                "Paste hosts when the production apps are deployed.",
            ],
        ],
        [16, 70, 70, 70, 70],
    )
    for r in range(2, env_ws.max_row + 1):
        env_ws.row_dimensions[r].height = 40

    how = wb.create_sheet("How this sheet is filled")
    write_rows(
        how,
        ["Sheet", "What it contains", "How it is used"],
        [
            ["Payloads", "SAPI row, PAPI/Exp transformed row, field map, current SQL", "Source of truth for mapping and queries"],
            ["Endpoints", "Paths plus query-param summary", "Quick lookup per resource"],
            ["Current field map", "One row per output field", "Use when changing a DWL"],
            ["Query parameters", "Exp vs PAPI vs SAPI params and headers", "Use when building Postman or RAML"],
            ["Authentication", "OIDC, PAPI hardcoded header, SAPI none", "Use when calling each layer"],
            ["Environments", "local / sandbox / presandbox / production hosts", "Import matching Postman environment"],
        ],
        [22, 70, 60],
    )
    for r in range(2, how.max_row + 1):
        how.row_dimensions[r].height = 28

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def q(key, value, desc=""):
    item = {"key": key, "value": value}
    if desc:
        item["description"] = desc
    return item


def header(key, value):
    return {"key": key, "value": value}


def url(raw, host_var, path, query):
    return {
        "raw": raw,
        "host": [host_var],
        "path": path,
        "query": query,
    }


def get_item(name, raw, host_var, path, query, headers, description):
    return {
        "name": name,
        "request": {
            "auth": {"type": "noauth"},
            "method": "GET",
            "header": headers,
            "url": url(raw, host_var, path, query),
            "description": description,
        },
    }


EXP_DATE = [
    q("fromDate", "{{fromDate}}"),
    q("toDate", "{{toDate}}"),
    q("pageNumber", "{{pageNumber}}"),
]
PAPI_DATE = [
    q("database", "{{database}}"),
    q("supplier", "{{supplier}}"),
    q("fromDate", "{{fromDate}}"),
    q("toDate", "{{toDate}}"),
    q("pageNumber", "{{pageNumber}}"),
]
PAPI_NODATE = [
    q("database", "{{database}}"),
    q("supplier", "{{supplier}}"),
    q("pageNumber", "{{pageNumber}}"),
]


def build_query_string(params):
    return "&".join(f"{p['key']}={p['value']}" for p in params)


def layer_gets(layer):
    items = []
    for r in RESOURCES:
        if layer == "exp":
            params = [] if r["exp_params"] == "(none)" else (
                [q("pageNumber", "{{pageNumber}}")] if r["exp_params"] == "pageNumber" else EXP_DATE
            )
            raw = "{{exp_base}}" + r["exp_path"]
            if params:
                raw += "?" + build_query_string(params)
            items.append(
                get_item(
                    f"GET {r['exp_path']}",
                    raw,
                    "{{exp_base}}",
                    [p for p in r["exp_path"].strip("/").split("/")],
                    params,
                    [header("Authorization", "Bearer {{access_token}}")],
                    f"{r['name']}. {r['exp_params']}. Bearer token only. Do not send database or supplier.",
                )
            )
        elif layer == "papi":
            params = PAPI_DATE if "fromDate" in r["papi_params"] else PAPI_NODATE
            raw = "{{papi_base}}" + r["papi_path"] + "?" + build_query_string(params)
            items.append(
                get_item(
                    f"GET {r['papi_path']}",
                    raw,
                    "{{papi_base}}",
                    [p for p in r["papi_path"].strip("/").split("/")],
                    params,
                    [header("Authorization", "Test")],
                    f"{r['name']}. {r['papi_params']}. Authorization is hardcoded Test. No variable.",
                )
            )
        else:
            params = PAPI_DATE if "fromDate" in r["sapi_params"] else PAPI_NODATE
            raw = "{{sapi_base}}" + r["sapi_path"] + "?" + build_query_string(params)
            items.append(
                get_item(
                    f"GET {r['sapi_path']}",
                    raw,
                    "{{sapi_base}}",
                    [p for p in r["sapi_path"].strip("/").split("/")],
                    params,
                    [],
                    f"{r['name']}. {r['sapi_params']}. No Authorization header.",
                )
            )
    return items


def token_request():
    return {
        "name": "Get Azure Token",
        "event": [
            {
                "listen": "test",
                "script": {
                    "type": "text/javascript",
                    "exec": [
                        "var json = pm.response.json();",
                        "if (json.access_token) {",
                        "    pm.collectionVariables.set('access_token', json.access_token);",
                        "}",
                    ],
                },
            }
        ],
        "request": {
            "auth": {"type": "noauth"},
            "method": "POST",
            "header": [header("Content-Type", "application/x-www-form-urlencoded")],
            "body": {
                "mode": "urlencoded",
                "urlencoded": [
                    {"key": "grant_type", "value": "client_credentials"},
                    {"key": "client_id", "value": "{{client_id}}"},
                    {"key": "client_secret", "value": "{{client_secret}}"},
                    {"key": "scope", "value": "{{scope}}"},
                ],
            },
            "url": {"raw": "{{token_url}}", "host": ["{{token_url}}"]},
            "description": "Kenya app SAT_Datashare_APP_API_SGKE_Reckitt. Put the secret in collection variable client_secret. Saves access_token for Exp GETs only.",
        },
    }


def env_file(name, env_id, exp_base, papi_base, sapi_base):
    return {
        "id": env_id,
        "name": name,
        "values": [
            {"key": "exp_base", "value": exp_base, "enabled": True},
            {"key": "papi_base", "value": papi_base, "enabled": True},
            {"key": "sapi_base", "value": sapi_base, "enabled": True},
            {"key": "database", "value": DATABASE, "enabled": True},
            {"key": "supplier", "value": SUPPLIER, "enabled": True},
            {"key": "fromDate", "value": "2026-08-13", "enabled": True},
            {"key": "toDate", "value": "2026-08-14", "enabled": True},
            {"key": "pageNumber", "value": "0", "enabled": True},
        ],
        "_postman_variable_scope": "environment",
    }


def build_collection():
    return {
        "info": {
            "_postman_id": "sat-datashare-exp-papi-sapi-20260921",
            "name": "SAT Datashare — Exp / PAPI / SAPI",
            "description": (
                "Eight datashare GETs for Exp, PAPI, and SAPI.\n\n"
                "Select a Postman environment: Local, Sandbox, Presandbox, or Production.\n"
                "That sets exp_base, papi_base, and sapi_base.\n\n"
                "Exp: run Get Azure Token first. GETs send only Authorization: Bearer {{access_token}}.\n"
                "Do not send database or supplier on Exp.\n"
                "Invoice / Order / Customer: fromDate, toDate, pageNumber.\n"
                "Stock: pageNumber only.\n"
                "Product / Route / SalesRep / Warehouse: no query params.\n\n"
                "PAPI: Authorization: Test (hardcoded, not a variable). Always send database and supplier.\n"
                "SAPI: no Authorization header. Always send database and supplier.\n\n"
                "Kenya defaults: database=sat_nobleoutlook, supplier=RECKITT BENCKISER.\n"
                "Dates are YYYY-MM-DD. pageNumber is 0-based."
            ),
            "schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json",
        },
        "variable": [
            {"key": "token_url", "value": "https://login.microsoftonline.com/f4faf003-d90b-4832-9df8-c8a22d29bdd4/oauth2/v2.0/token"},
            {"key": "client_id", "value": "de772600-0d1f-4492-90b0-a57fcc284cf9"},
            {"key": "client_secret", "value": ""},
            {"key": "scope", "value": "https://graph.microsoft.com/.default"},
            {"key": "access_token", "value": ""},
            {"key": "database", "value": DATABASE},
            {"key": "supplier", "value": SUPPLIER},
            {"key": "fromDate", "value": "2026-08-13"},
            {"key": "toDate", "value": "2026-08-14"},
            {"key": "pageNumber", "value": "0"},
        ],
        "item": [
            {
                "name": "Auth",
                "description": "Azure token for Exp only. PAPI and SAPI do not use this token.",
                "item": [token_request()],
            },
            {
                "name": "Exp",
                "description": "Experience API. Bearer token only. No extra headers. No database/supplier.",
                "item": layer_gets("exp"),
            },
            {
                "name": "PAPI",
                "description": "Process API. Authorization hardcoded to Test. No other headers.",
                "item": layer_gets("papi"),
            },
            {
                "name": "SAPI",
                "description": "System API. No Authorization header. database and supplier required.",
                "item": layer_gets("sapi"),
            },
        ],
    }


def main():
    xlsx = MAPPING / "SAT-Datashare-SAPI-PAPI-EAPI-payloads.xlsx"
    xlsx_alias = POSTMAN / "SAT-Datashare-SAPI-PAPI-EAPI-payloads (2) (1).xlsx"
    build_excel(xlsx)
    build_excel(xlsx_alias)

    collection = build_collection()
    (POSTMAN / "SAT-Datashare-Exp-PAPI-SAPI.postman_collection.json").write_text(
        json.dumps(collection, indent=2) + "\n", encoding="utf-8"
    )
    (POSTMAN / "SAT-Datashare-8-GETs.postman_collection.json").write_text(
        json.dumps(collection, indent=2) + "\n", encoding="utf-8"
    )

    envs = [
        (
            "SAT-Datashare-local.postman_environment.json",
            env_file(
                "SAT Datashare — Local",
                "sat-datashare-env-local",
                "http://localhost:8081",
                "http://localhost:8081",
                "http://localhost:8081",
            ),
        ),
        (
            "SAT-Datashare-sandbox.postman_environment.json",
            env_file(
                "SAT Datashare — Sandbox",
                "sat-datashare-env-sandbox",
                "https://exp-sat-datashare-sand-api-5nct48.2ky31l-2.deu-c1.eu1.cloudhub.io",
                "https://prc-solutech-api-v1-5nct48.2ky31l-1.deu-c1.eu1.cloudhub.io",
                "https://sys-solutech-api-v1-5nct48.2ky31l-2.deu-c1.eu1.cloudhub.io",
            ),
        ),
        (
            "SAT-Datashare-presandbox.postman_environment.json",
            env_file(
                "SAT Datashare — Presandbox",
                "sat-datashare-env-presandbox",
                "",
                "",
                "",
            ),
        ),
        (
            "SAT-Datashare-production.postman_environment.json",
            env_file(
                "SAT Datashare — Production",
                "sat-datashare-env-production",
                "",
                "",
                "",
            ),
        ),
    ]
    for name, body in envs:
        (POSTMAN / name).write_text(json.dumps(body, indent=2) + "\n", encoding="utf-8")

    # Keep older env names pointing at the same local/sandbox values.
    (POSTMAN / "SAT-Datashare-localhost.postman_environment.json").write_text(
        json.dumps(envs[0][1], indent=2) + "\n", encoding="utf-8"
    )
    (POSTMAN / "SAT-Datashare-CloudHub.postman_environment.json").write_text(
        json.dumps(envs[1][1], indent=2) + "\n", encoding="utf-8"
    )
    print("wrote", xlsx)
    print("wrote", xlsx_alias)
    print("wrote collection and 4 environments")


if __name__ == "__main__":
    main()
